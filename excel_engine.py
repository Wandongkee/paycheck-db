"""Workbook conversion and header-based payroll processing (no Streamlit dependency)."""
from copy import copy
import datetime as dt
import io
import math
import re
import shutil
import subprocess
import tempfile
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import from_excel


class InputError(ValueError):
    pass


ALIASES = {
    'name': ('성명', '이름', '사원명', '직원명'),
    'hire': ('입사일', '입사일자', '입사년월일', '입사일(년월일)'),
    'dept': ('부서', '부서명', '본부', '본부명', '소속'),
    'job': ('직종', '직종명', '직책'),
    'base': ('지급총액',), 'subtract1': ('식대',), 'subtract2': ('연차수당',),
    'amount': ('OT금액', 'OT수당', '시간외수당합계', '연장근로수당합계', '계'),
    'early': ('조출점심저녁', '조출점심저녁시간'),
    'extension': ('연장OT', '연장OT시간', '연장시간'),
    'night': ('야간OT', '야간OT시간', '야간시간'),
    'holiday_days': ('휴일근무', '휴일근무일수'),
    'holiday_hours': ('휴일OT', '휴일OT시간'),
}


def normalize(value):
    return re.sub(r'[\s_\-./():（）]+', '', str(value or '')).casefold()


def to_xlsx(data, filename, executable=None):
    """Use Calc's workbook converter, never a values-only DataFrame export."""
    if data.startswith(b'PK') and zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if 'xl/workbook.xml' not in archive.namelist():
                raise InputError('Excel 통합문서가 아닙니다.')
        return data  # Some ERP exports have an incorrect .xls extension.
    if Path(filename).suffix.lower() != '.xls':
        raise InputError('파일 내용을 읽을 수 없습니다. 정상적인 .xls/.xlsx 파일인지 확인하세요.')
    executable = executable or shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        raise InputError('XLS 변환기가 설치되지 않았습니다. 서버에 packages.txt의 libreoffice-calc 설치가 필요합니다.')
    with tempfile.TemporaryDirectory(prefix='payroll-') as folder:
        root = Path(folder)
        profile = root / 'profile'
        user = profile / 'user'
        user.mkdir(parents=True)
        # Separate profile for each conversion; do not execute workbook macros.
        (user / 'registrymodifications.xcu').write_text(
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
            '</item></oor:items>', encoding='utf-8')
        source = root / 'input.xls'
        source.write_bytes(data)
        try:
            result = subprocess.run(
                [str(executable), '-env:UserInstallation=' + profile.as_uri(),
                 '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
                 '--outdir', str(root), str(source)],
                capture_output=True, timeout=90, check=False)
        except subprocess.TimeoutExpired as exc:
            raise InputError('XLS 변환 시간이 초과되었습니다. 파일을 확인해주세요.') from exc
        output = root / 'input.xlsx'
        if result.returncode or not output.exists():
            raise InputError('XLS 자동 변환에 실패했습니다. 암호 설정 또는 파일 손상 여부를 확인해주세요.')
        converted = output.read_bytes()
        try:
            wb = openpyxl.load_workbook(io.BytesIO(converted), read_only=True)
            wb.close()
        except Exception as exc:
            raise InputError('변환 결과가 정상적인 XLSX 파일이 아닙니다.') from exc
        return converted


def read_book(data, values=False):
    return openpyxl.load_workbook(io.BytesIO(data), data_only=values)


def headers(ws, row, end_row=None):
    end_row = end_row or row
    labels = {}
    for col in range(1, ws.max_column + 1):
        parts = []
        for r in range(row, end_row + 1):
            value = ws.cell(r, col).value
            if value is None:
                for merged in ws.merged_cells.ranges:
                    if merged.min_row <= r <= merged.max_row and merged.min_col <= col <= merged.max_col:
                        value = ws.cell(merged.min_row, merged.min_col).value
                        break
            # ERP 'values only' copies may lose merged-cell definitions.
            # Borrow only recognized OT group labels, never an arbitrary neighboring header.
            if value is None and col > 1:
                left = ws.cell(r, col - 1).value
                if normalize(left) in {normalize(v) for v in ('조출, 점심, 저녁O/T', '조출, 저녁O/T', '조출, 점심O/T', '기본연장O/T', '야간O/T', '휴일근무', '휴일O/T')}:
                    value = left
            if value is not None and str(value).strip() and str(value).strip() not in parts:
                parts.append(str(value).strip())
        if parts:
            labels[col] = ' | '.join(parts)
    return labels


def candidates(labels, role):
    aliases = {normalize(x) for x in ALIASES.get(role, ())}
    result = []
    time_groups = {'early': ('조출점심저녁ot', '조출저녁ot', '조출점심ot'),
                   'extension': ('기본연장ot',), 'night': ('야간ot',),
                   'holiday_days': ('휴일근무',), 'holiday_hours': ('휴일ot',)}
    for col, value in labels.items():
        parts = [normalize(x).replace(',', '') for x in value.split(' | ')]
        if len(parts) == 1 and parts[0] in aliases:
            result.append(col)
        elif len(parts) > 1:
            if role == 'amount' and '계' in parts:
                result.append(col)
            elif role in time_groups:
                if any(x in parts for x in time_groups[role]) and any(x in parts for x in ('시간', '일', '시')) and '금액' not in parts:
                    result.append(col)
            elif role not in ('amount',) and any(x in aliases for x in parts):
                result.append(col)
    return result


def detect_header_end(ws, start):
    labels = headers(ws, start)
    names, hires = candidates(labels, 'name'), candidates(labels, 'hire')
    if len(names) == 1 and len(hires) == 1:
        for row in range(start + 1, min(ws.max_row, start + 15) + 1):
            try:
                person_key(ws.cell(row, names[0]).value, ws.cell(row, hires[0]).value, ws.parent.epoch)
                return row - 1
            except (InputError, ValueError, TypeError, OverflowError):
                continue
    return start


def detect_header(ws):
    scores = []
    for row in range(1, min(ws.max_row, 40) + 1):
        labels = {c.column: str(c.value).strip() for c in ws[row] if c.value is not None}
        score = (100 * bool(candidates(labels, 'name')) + 50 * bool(candidates(labels, 'hire'))
                 + sum(bool(candidates(labels, role)) for role in ALIASES))
        scores.append((score, -row))
    return -max(scores)[1] if scores else 1


def number(value, context, blank_zero=True):
    if value is None or str(value).strip() in ('', '-'):
        if blank_zero:
            return Decimal(0)
        raise InputError(f'{context}: 금액 또는 수식 계산 결과가 없습니다. Excel에서 재계산 후 저장해주세요.')
    try:
        result = Decimal(str(value).replace(',', '').strip())
        if not result.is_finite():
            raise InvalidOperation
        return result
    except (InvalidOperation, ValueError):
        raise InputError(f'{context}: 숫자로 읽을 수 없는 값입니다 ({value}).') from None


def date_key(value, epoch):
    if value is None or str(value).strip() == '':
        raise InputError('입사일자가 비어 있습니다.')
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise InputError('입사일자가 올바르지 않습니다.')
        if 19000101 <= value <= 99991231 and value == int(value):
            value = str(int(value))
        else:
            value = from_excel(value, epoch)
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime('%Y%m%d')
    text = str(value).strip()
    match = re.fullmatch(r'(\d{4})\D+(\d{1,2})\D+(\d{1,2})(?:[ T]00:00:00)?\D*', text)
    try:
        if match:
            return dt.date(*map(int, match.groups())).strftime('%Y%m%d')
        return dt.datetime.strptime(text, '%Y%m%d').strftime('%Y%m%d')
    except ValueError:
        raise InputError(f'입사일자를 해석할 수 없습니다: {value}') from None


def person_key(name, hire, epoch):
    name = str(name or '').split('(')[0].strip()
    if not name:
        raise InputError('이름이 비어 있습니다.')
    return name, date_key(hire, epoch)


def validate_mapping(ws, mapping, roles):
    cols = [mapping.get(role) for role in roles]
    if any(not isinstance(col, int) or not 1 <= col <= ws.max_column for col in cols):
        raise InputError(f'{ws.title}: 필요한 열을 모두 선택해주세요.')
    if len(cols) != len(set(cols)):
        raise InputError(f'{ws.title}: 서로 다른 항목에 같은 열을 선택할 수 없습니다.')


def value_at(ws, cached, row, col, context):
    cell = ws.cell(row, col)
    value = cached.cell(row, col).value
    return number(value, context, blank_zero=cell.data_type != 'f')


def load_ot(sources):
    result = {}
    for data, sheet, header, mapping, group, filename in sources:
        wb, cached = read_book(data), read_book(data, True)
        ws, values = wb[sheet], cached[sheet]
        validate_mapping(ws, mapping, ('name', 'hire', 'amount'))
        for row in range(header + 1, ws.max_row + 1):
            name = values.cell(row, mapping['name']).value
            if not name or normalize(name) in ('합계', '총계', '소계'):
                continue
            where = f'{filename} / {sheet} / {row}행'
            try:
                key = (group,) + person_key(name, values.cell(row, mapping['hire']).value, wb.epoch)
                amount = number(values.cell(row, mapping['amount']).value, where, blank_zero=False)
                if key in result:
                    raise InputError('동일 본부의 이름·입사일자가 중복됩니다. OT 원본을 확인해주세요.')
                result[key] = amount
            except InputError as exc:
                raise InputError(f'{where}: {exc}') from exc
    return result


SALARY_ROLES = ('dept', 'name', 'hire', 'job', 'ot1', 'ot2', 'ot3', 'ot4', 'base', 'subtract1', 'subtract2')


def integrate(data, sheet, header, mapping, ot_sources):
    wb, cached = read_book(data), read_book(data, True)
    ws, values = wb[sheet], cached[sheet]
    validate_mapping(ws, mapping, SALARY_ROLES)
    ot = load_ot(ot_sources)
    supplied_groups = {source[4] for source in ot_sources}
    out = wb.create_sheet('대조내역')
    out.append(['원본시트', '원본행', '부서', '이름', '입사일자', '직종',
                '급여DB OT 합산', 'OT 파일 금액', '금액 일치', '차액', '확인사항'])
    matched = 0
    for row in range(header + 1, ws.max_row + 1):
        name = values.cell(row, mapping['name']).value
        if not name or normalize(name) in ('합계', '총계', '소계'):
            continue
        if any(normalize(values.cell(row, col).value) in ('부서별총계', '사업장별총계', '총계')
               for col in range(1, min(ws.max_column, 4) + 1)):
            continue
        dept = str(values.cell(row, mapping['dept']).value or '').strip()
        job = str(values.cell(row, mapping['job']).value or '').strip()
        hire = values.cell(row, mapping['hire']).value
        total, found, equal, difference, status = None, None, None, None, ''
        try:
            if job == '양중(T/C)':
                key = person_key(name, hire, wb.epoch)
                total = sum(value_at(ws, values, row, mapping[f'ot{i}'], f'{sheet} {row}행 OT{i}') for i in range(1, 5))
                normalized_dept = normalize(dept)
                group = '운영1' if '운영1' in normalized_dept else '운영2' if '운영2' in normalized_dept else '운영' if '운영' in normalized_dept else None
                if group is None:
                    status = '부서 확인 필요: 운영1/운영2/운영 구분 불가'
                else:
                    found = ot.get((group,) + key)
                    if found is None:
                        status = 'OT 파일 미제공' if group not in supplied_groups else 'OT 미매칭'
                    else:
                        equal = total == found
                        matched += 1
                        status = '일치' if equal else '금액 불일치'
            else:
                status = 'OT 매칭 대상 외'
            difference = (value_at(ws, values, row, mapping['base'], f'{sheet} {row}행 기준금액')
                          - value_at(ws, values, row, mapping['subtract1'], f'{sheet} {row}행 차감1')
                          - (total or Decimal(0))
                          - value_at(ws, values, row, mapping['subtract2'], f'{sheet} {row}행 차감2'))
        except InputError as exc:
            status = str(exc)
        out.append([sheet, row, dept, str(name), str(hire or ''), job,
                    float(total) if total is not None else None,
                    float(found) if found is not None else None, equal,
                    float(difference) if difference is not None else None, status])
    out.freeze_panes = 'G2'
    out.auto_filter.ref = out.dimensions
    for cell in out[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='24527A')
    for col in 'ABCDEFGHIJ':
        out.column_dimensions[col].width = 20
    out.column_dimensions['K'].width = 65
    for row in out.iter_rows(min_row=2, min_col=7, max_col=10):
        for cell in row:
            if cell.column != 9:
                cell.number_format = '#,##0'
    rows = list(out.values)
    build_fixed_output(wb, ws, header, mapping, rows)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue(), matched, rows


def statement(entries):
    if not entries:
        raise InputError('처리할 시트를 선택해주세요.')
    data = entries[0][0]
    wb, cached = read_book(data), read_book(data, True)
    roles = ('extension', 'night', 'holiday_days', 'holiday_hours', 'early')
    labels = ('연장OT', '야간OT', '휴일근무', '휴일OT', '조출점심저녁')
    units = ('H', 'H', 'D', 'H', 'H')
    for _, sheet, header, mapping, _, filename in entries:
        ws, values = wb[sheet], cached[sheet]
        validate_mapping(ws, mapping, ('name',) + roles)
        start = 53  # Keep existing downstream BA:BE coordinates.
        if any(ws.cell(r, c).value is not None for r in range(header, ws.max_row + 1) for c in range(53, 58)):
            raise InputError(f'{filename}/{sheet}: BA~BE열에 기존 값이 있어 덮어쓰지 않습니다.')
        for i, label in enumerate(labels):
            ws.cell(header, start + i, label + ' 문구')
            ws.column_dimensions[openpyxl.utils.get_column_letter(start + i)].width = 22
        for row in range(header + 1, ws.max_row + 1):
            name = values.cell(row, mapping['name']).value
            if not name or normalize(name) in ('합계', '총계', '소계'):
                continue
            for i, (role, label, unit) in enumerate(zip(roles, labels, units)):
                value = value_at(ws, values, row, mapping[role], f'{filename}/{sheet}/{row}행 {label}')
                text = format(value, 'f')
                if '.' in text:
                    text = text.rstrip('0').rstrip('.')
                ws.cell(row, start + i, f'{label}:{text}{unit}')
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def vlookup_values(data):
    wb, cached = read_book(data), read_book(data, True)
    for ws in wb:
        for row in ws:
            for cell in row:
                if cell.data_type == 'f' and 'VLOOKUP' in cell.value.upper():
                    value = cached[ws.title][cell.coordinate].value
                    if value is None:
                        raise InputError(f'{ws.title}!{cell.coordinate}: 계산 결과가 없습니다. Excel에서 재계산 후 저장해주세요.')
                    cell.value = value
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# Confirmed payroll headers. No positional fallback when labels are missing.
ALIASES.update({
    'ot1': ('조출점심저녁OT',),
    'ot2': ('연장수당',),
    'ot3': ('야간수당',),
    'ot4': ('휴일수당(1)',),
})


def build_fixed_output(wb, source, header, mapping, audit_rows):
    """Stable downstream coordinates; source workbook sheets remain intact."""
    result = wb.create_sheet('급여통합결과', 0)
    pinned = {'dept': 2, 'name': 4, 'hire': 6, 'job': 11,
              'ot1': 20, 'ot2': 17, 'ot3': 18, 'ot4': 19,
              'subtract1': 24, 'subtract2': 45, 'base': 47}
    destination = {mapping[role]: col for role, col in pinned.items()}
    reserved = set(pinned.values()) | {21, 22, 23, 48}
    available = (c for c in range(1, source.max_column + 60) if c not in reserved)
    for col in range(1, source.max_column + 1):
        if col not in destination:
            destination[col] = next(available)
    source_title = "'" + source.title.replace("'", "''") + "'"
    source_rows = [header] + [row[1] for row in audit_rows[1:]]
    for new_row, old_row in enumerate(source_rows, 1):
        result.row_dimensions[new_row].height = source.row_dimensions[old_row].height
        for old_col, new_col in destination.items():
            cell = source.cell(old_row, old_col)
            target = result.cell(new_row, new_col)
            target.value = f'={source_title}!{cell.coordinate}' if cell.data_type == 'f' else cell.value
            if cell.has_style:
                target._style = copy(cell._style)
            if cell.comment:
                target.comment = copy(cell.comment)
        if new_row > 1:
            result.cell(new_row, 4, str(result.cell(new_row, 4).value or '').split('(')[0].strip())
    for old_col, new_col in destination.items():
        old_letter = openpyxl.utils.get_column_letter(old_col)
        new_letter = openpyxl.utils.get_column_letter(new_col)
        result.column_dimensions[new_letter].width = source.column_dimensions[old_letter].width
    for col, title in ((21, '급여DB OT 합산'), (22, 'OT 파일 금액'), (23, '금액 일치'), (48, '차액')):
        result.cell(1, col, title)
        result.cell(1, col).font = Font(bold=True)
        result.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20
    status_col = max(48, max(destination.values())) + 1
    result.cell(1, status_col, '확인사항')
    result.column_dimensions[openpyxl.utils.get_column_letter(status_col)].width = 55
    for row, audit in enumerate(audit_rows[1:], 2):
        if audit[6] is not None:
            result.cell(row, 21, f'=SUM(Q{row}:T{row})')
        if audit[7] is not None:
            result.cell(row, 22, audit[7])
        if audit[6] is not None:
            result.cell(row, 23, f'=IF(V{row}="","미매칭",U{row}=V{row})')
        if audit[9] is not None:
            result.cell(row, 48, f'=AU{row}-X{row}-U{row}-AS{row}')
        result.cell(row, status_col, audit[10])
        for col in (21, 22, 48):
            result.cell(row, col).number_format = '#,##0'
    result.freeze_panes = 'H2'
    result.auto_filter.ref = result.dimensions
    wb.active = 0
    wb.calculation = openpyxl.workbook.properties.CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)


NOTICE_ROLES = ('extension', 'night', 'holiday_days', 'holiday_hours', 'early')
NOTICE_LABELS = ('연장OT', '야간OT', '휴일근무', '휴일OT', '조출점심저녁')
NOTICE_UNITS = ('H', 'H', 'D', 'H', 'H')


def notice_texts(values):
    result = []
    for value, label, unit in zip(values, NOTICE_LABELS, NOTICE_UNITS):
        text = format(value, 'f')
        if '.' in text:
            text = text.rstrip('0').rstrip('.')
        result.append(f'{label}:{text}{unit}')
    return result


def legacy_statement(data):
    """Section 3: original all-sheet BA:BE text operation, without mapping UI."""
    wb, cached = read_book(data), read_book(data, True)
    for ws in wb:
        for col in range(10, 53):
            dimension = ws.column_dimensions[openpyxl.utils.get_column_letter(col)]
            dimension.outlineLevel = 1
            dimension.hidden = True
        for col in range(53, 58):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
        for row in range(8, ws.max_row + 1):
            if not ws.cell(row, 5).value:
                continue
            numbers = []
            for col in (12, 14, 16, 18, 10):
                try:
                    numbers.append(number(cached[ws.title].cell(row, col).value, 'OT 시간'))
                except InputError:
                    numbers.append(Decimal(0))
            for col, text in enumerate(notice_texts(numbers), 53):
                ws.cell(row, col, text)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()



def append_ot_notices(sources):
    """Append texts to selected OT sheets, preserving every existing cell/cache."""
    grouped = {}
    for data, sheet, header, mapping, group, filename in sources:
        grouped.setdefault((group, filename, data), []).append((sheet, header, mapping))
    outputs = []
    for (group, filename, data), sheets in grouped.items():
        wb, cached = read_book(data), read_book(data, True)
        updates = {}
        for sheet, header, mapping in sheets:
            ws, values = wb[sheet], cached[sheet]
            validate_mapping(ws, mapping, ('name',) + NOTICE_ROLES)
            start = ws.max_column + 1
            texts = {header: [label + ' 안내' for label in NOTICE_LABELS]}
            for row in range(header + 1, ws.max_row + 1):
                name = values.cell(row, mapping['name']).value
                if not name or normalize(name) in ('합계', '총계', '소계'):
                    continue
                numbers = [value_at(ws, values, row, mapping[role],
                           f'{filename}/{sheet}/{row}행 {label}')
                           for role, label in zip(NOTICE_ROLES, NOTICE_LABELS)]
                texts[row] = notice_texts(numbers)
            updates[sheet] = (start, texts)
        outputs.append((group, filename, append_text_cells(data, updates)))
    return outputs


def append_text_cells(data, updates):
    # Patch only the selected sheet XML. Re-saving an entire workbook would
    # discard formula caches, and can change unrelated sheets or Excel features.
    import posixpath
    from xml.dom import minidom
    main_ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rel_ns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(io.BytesIO(data)) as original:
        book = minidom.parseString(original.read('xl/workbook.xml'))
        rels = minidom.parseString(original.read('xl/_rels/workbook.xml.rels'))
        targets = {r.getAttribute('Id'): r.getAttribute('Target')
                   for r in rels.getElementsByTagName('Relationship')}
        paths = {}
        for sheet in book.getElementsByTagNameNS(main_ns, 'sheet'):
            target = targets[sheet.getAttributeNS(rel_ns, 'id')]
            paths[sheet.getAttribute('name')] = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        replacements = {}
        for name, (start, texts) in updates.items():
            path = paths[name]
            doc = minidom.parseString(original.read(path))
            root = doc.documentElement
            prefix = root.prefix + ':' if root.prefix else ''
            def element(tag, attributes=None):
                node = doc.createElementNS(main_ns, prefix + tag)
                for key, value in (attributes or {}).items():
                    node.setAttribute(key, str(value))
                return node
            sheet_data = doc.getElementsByTagNameNS(main_ns, 'sheetData')[0]
            rows = {int(n.getAttribute('r')): n for n in sheet_data.childNodes
                    if n.nodeType == n.ELEMENT_NODE and n.localName == 'row'}
            for number, values in sorted(texts.items()):
                row = rows.get(number)
                if row is None:
                    row = element('row', {'r': number})
                    next_row = next((rows[r] for r in sorted(rows) if r > number), None)
                    sheet_data.insertBefore(row, next_row)
                    rows[number] = row
                if row.hasAttribute('spans'):
                    row.removeAttribute('spans')
                for col, value in enumerate(values, start):
                    address = f'{openpyxl.utils.get_column_letter(col)}{number}'
                    cell = element('c', {'r': address, 't': 'inlineStr'})
                    inline, text = element('is'), element('t')
                    text.appendChild(doc.createTextNode(value))
                    inline.appendChild(text)
                    cell.appendChild(inline)
                    row.appendChild(cell)
            dims = doc.getElementsByTagNameNS(main_ns, 'dimension')
            if dims:
                min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(dims[0].getAttribute('ref'))
                end_col = openpyxl.utils.get_column_letter(max(max_col, start + 4))
                dims[0].setAttribute('ref', f'{openpyxl.utils.get_column_letter(min_col)}{min_row}:{end_col}{max(max_row, max(texts))}')
            columns = doc.getElementsByTagNameNS(main_ns, 'cols')
            if columns:
                columns = columns[0]
            else:
                columns = element('cols')
                root.insertBefore(columns, sheet_data)
            columns.appendChild(element('col', {'min': start, 'max': start + 4, 'width': 23, 'customWidth': 1}))
            replacements[path] = doc.toxml(encoding='utf-8')
        output = io.BytesIO()
        with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as result:
            for info in original.infolist():
                result.writestr(info, replacements.get(info.filename, original.read(info.filename)))
        return output.getvalue()
