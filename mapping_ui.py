import streamlit as st
from openpyxl.utils import get_column_letter
from excel_engine import candidates, detect_header, detect_header_end, headers

LABELS = {
    'dept': '부서', 'name': '이름', 'hire': '입사일자', 'job': '직종',
    'ot1': '조출점심저녁OT', 'ot2': '연장수당', 'ot3': '야간수당', 'ot4': '휴일수당(1)',
    'base': '차액 기준금액 (기존 결과 AU열)',
    'subtract1': '차감금액 1 (기존 결과 X열)', 'subtract2': '차감금액 2 (기존 결과 AS열)',
    'amount': 'OT 총금액 (기존 T열)', 'early': '조출점심저녁 시간',
    'extension': '연장OT 시간', 'night': '야간OT 시간',
    'holiday_days': '휴일근무 일수', 'holiday_hours': '휴일OT 시간',
}

def mapping_panel(ws, roles, key):
    header = int(st.number_input('열 제목이 있는 행', min_value=1, max_value=max(1, ws.max_row),
                                value=detect_header(ws), key=key + ':header'))
    end = int(st.number_input('제목 마지막 행 (다음 행부터 데이터)', min_value=header,
        max_value=max(header, ws.max_row), value=detect_header_end(ws, header), key=key + ':end:' + str(header)))
    labels = headers(ws, header, end)
    options = [None] + list(labels)
    mapping = {}
    for role in roles:
        matches = candidates(labels, role)
        default = matches[0] if len(matches) == 1 else None
        mapping[role] = st.selectbox(LABELS[role], options, index=options.index(default),
            format_func=lambda col: '선택해주세요' if col is None else f'{labels[col]} ({get_column_letter(col)}열)',
            key=f'{key}:{header}:{role}')
    selected = [v for v in mapping.values() if v is not None]
    ready = len(selected) == len(roles) and len(selected) == len(set(selected))
    if not ready:
        st.info('필요한 항목을 모두 선택해주세요. 서로 다른 항목에는 서로 다른 열을 선택해야 합니다.')
    return end, mapping, ready
