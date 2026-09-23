import datetime as dt
import io
import shutil
import zipfile
from unittest.mock import patch
import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.datetime import WINDOWS_EPOCH
from excel_engine import (InputError, SALARY_ROLES, candidates, date_key, detect_header,
    headers, integrate, read_book, statement, to_xlsx, vlookup_values)


def save(wb):
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def fixture(reverse=False):
    labels = ['부서', '성명', '입사일자', '직종', '조출점심저녁OT', '연장수당',
              '야간수당', '휴일수당(1)', '기준금액', '차감1', '차감2']
    vals = ['운영1본부', '홍길동(기사)', dt.datetime(2020, 1, 2), '양중(T/C)',
            100, 200, 300, 400, 5000, 500, 100]
    pairs = list(zip(SALARY_ROLES, labels, vals))
    if reverse:
        pairs.reverse()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '급여'
    ws.merge_cells('A1:C1')
    ws['A1'] = '급여대장'
    ws.append([x[1] for x in pairs])
    ws.append([x[2] for x in pairs])
    ws['A2'].font = Font(bold=True)
    wb.create_sheet('안내')['A1'] = '=1+2'
    mapping = {role: i + 1 for i, (role, _, _) in enumerate(pairs)}
    ot = openpyxl.Workbook()
    ot.active.title = 'OT'
    ot.active.append(['OT금액', '입사일', '이름', '연장OT시간', '야간OT시간', '휴일근무일수', '휴일OT시간', '조출점심저녁시간'])
    ot.active.append([1000, '2020.1.2', '홍길동', 16, 8, 4, 9, 0])
    sources = [(save(ot), 'OT', 1, {'amount': 1, 'hire': 2, 'name': 3}, '운영1', 'ot.xlsx')]
    return save(wb), mapping, sources


@pytest.mark.parametrize('reverse', [False, True])
def test_reordered_columns_and_original_preserved(reverse):
    data, mapping, sources = fixture(reverse)
    result, count, rows = integrate(data, '급여', 2, mapping, sources)
    assert count == 1
    fixed = read_book(result).active
    assert fixed['U2'].value == '=SUM(Q2:T2)'
    assert fixed['V2'].value == 1000
    assert fixed['W2'].value == '=IF(V2="","미매칭",U2=V2)'
    assert fixed['AV2'].value == '=AU2-X2-U2-AS2'
    assert [fixed.cell(2, c).value for c in range(17, 21)] == [200, 300, 400, 100]
    assert fixed['AU2'].value == 5000
    assert rows[1][6:10] == (1000, 1000, True, 3400)
    before, after = read_book(data), read_book(result)
    for ws in before:
        assert list(ws.values) == list(after[ws.title].values)
        assert str(ws.merged_cells) == str(after[ws.title].merged_cells)
    assert after['급여']['A2'].font.bold
    labels = headers(after['급여'], 2)
    assert detect_header(after['급여']) == 2
    for role in ('name', 'hire', 'dept', 'job', 'ot1', 'ot2', 'ot3', 'ot4'):
        assert candidates(labels, role) == [mapping[role]]


def test_duplicate_ot_rejected():
    data, mapping, sources = fixture()
    with pytest.raises(InputError, match='중복'):
        integrate(data, '급여', 2, mapping, sources * 2)


def test_zero_ot_and_wrong_group():
    data, mapping, sources = fixture()
    wb = read_book(sources[0][0])
    wb.active['A2'] = 0
    sources[0] = (save(wb),) + sources[0][1:]
    _, count, rows = integrate(data, '급여', 2, mapping, sources)
    assert count == 1 and rows[1][7] == 0 and rows[1][10] == '금액 불일치'
    sources[0] = sources[0][:4] + ('운영2', 'ot.xlsx')
    _, count, rows = integrate(data, '급여', 2, mapping, sources)
    assert count == 0 and rows[1][10] == 'OT 파일 미제공'


def test_uncached_salary_formula_is_not_zero():
    data, mapping, sources = fixture()
    wb = read_book(data)
    wb['급여'].cell(3, mapping['ot1'], '=50+50')
    _, count, rows = integrate(save(wb), '급여', 2, mapping, sources)
    assert count == 0
    assert rows[1][6] is None and '계산 결과' in rows[1][10]


def test_mapping_missing_duplicate():
    data, mapping, sources = fixture()
    mapping['ot1'] = mapping['ot2']
    with pytest.raises(InputError, match='같은 열'):
        integrate(data, '급여', 2, mapping, sources)
    mapping['ot1'] = None
    with pytest.raises(InputError, match='모두 선택'):
        integrate(data, '급여', 2, mapping, sources)


@pytest.mark.parametrize('value', [dt.datetime(2020, 1, 2), '2020-1-2', '2020/01/02',
                                  '20200102', 20200102, 43832])
def test_dates(value):
    assert date_key(value, WINDOWS_EPOCH) == '20200102'


def test_invalid_date_and_duplicate_headers():
    with pytest.raises(InputError):
        date_key('2020-13-32', WINDOWS_EPOCH)
    assert candidates({1: '성명', 3: '성명'}, 'name') == [1, 3]


def test_statement_reordered_no_overwrite():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['야간OT', '성명', '조출점심저녁', '휴일OT', '연장OT', '휴일근무'])
    ws.append([2, '홍길동', 1.5, 3, 4, 1])
    mapping = dict(night=1, name=2, early=3, holiday_hours=4, extension=5, holiday_days=6)
    result = read_book(statement([(save(wb), ws.title, 1, mapping, '', 'ot.xlsx')]))
    assert list(result.active.values)[1][:6] == (2, '홍길동', 1.5, 3, 4, 1)
    assert list(result.active.values)[1][52:57] == ('연장OT:4H', '야간OT:2H', '휴일근무:1D', '휴일OT:3H', '조출점심저녁:1.5H')


def test_xlsx_mislabeled_as_xls_is_preserved():
    data, _, _ = fixture()
    assert to_xlsx(data, 'erp.xls') == data


def test_missing_converter_does_not_export_values_only():
    with patch('excel_engine.shutil.which', return_value=None):
        with pytest.raises(InputError, match='설치'):
            to_xlsx(b'legacy', 'erp.xls')


def test_vlookup_cache_missing_stops():
    wb = openpyxl.Workbook()
    wb.active['A1'] = '=VLOOKUP(1,B1:C2,2,FALSE)'
    with pytest.raises(InputError, match='계산 결과'):
        vlookup_values(save(wb))


def test_real_xls_conversion():
    executable = shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        pytest.skip('LibreOffice unavailable locally; required in Ubuntu CI')
    import xlwt
    wb = xlwt.Workbook()
    ws = wb.add_sheet('급여')
    style = xlwt.easyxf('font: bold on; pattern: pattern solid, fore_colour yellow;')
    ws.write_merge(0, 0, 0, 2, '급여대장', style)
    ws.col(0).width = 6000
    ws.write(1, 0, dt.datetime(2020, 1, 2), xlwt.easyxf(num_format_str='YYYY-MM-DD'))
    ws.write(1, 1, 100)
    ws.write(1, 2, xlwt.Formula('B2*2'))
    wb.add_sheet('두번째').write(0, 0, '보존')
    stream = io.BytesIO()
    wb.save(stream)
    result = to_xlsx(stream.getvalue(), 'erp.xls', executable)
    out = read_book(result)
    assert out.sheetnames == ['급여', '두번째']
    assert str(out['급여'].merged_cells) == 'A1:C1'
    assert out['급여']['A1'].font.bold
    assert out['급여']['A1'].fill.fgColor.rgb[-6:] == 'FFFF00'
    assert out['급여'].column_dimensions['A'].width > 20
    assert out['급여']['A2'].value == dt.datetime(2020, 1, 2)
    assert out['급여']['C2'].data_type == 'f'
    assert read_book(result, True)['급여']['C2'].value == 200


def test_app_initial_render():
    from streamlit.testing.v1 import AppTest
    app = AppTest.from_file(str(__import__('pathlib').Path('opp.py').resolve())).run(timeout=20)
    assert not app.exception
    assert all(button.disabled for button in app.button)


def test_app_upload_mapping_and_download():
    from streamlit.testing.v1 import AppTest
    from pathlib import Path
    data, _, sources = fixture(reverse=True)
    class Upload(io.BytesIO):
        def __init__(self, data, name):
            super().__init__(data)
            self.name = name
    db, ot = Upload(data, 'erp.xlsx'), Upload(sources[0][0], 'ot.xlsx')
    def uploads(label, **kwargs):
        if label.startswith('메인 급여DB'):
            return db
        if label == '운영1 OT 파일':
            return [ot]
        return [] if kwargs.get('accept_multiple_files') else None
    with patch('streamlit.file_uploader', side_effect=uploads):
        app = AppTest.from_file(str(Path('opp.py').resolve())).run(timeout=20)
        assert not app.exception
        required = {'차액 기준금액 (기존 결과 AU열)': 3,
                    '차감금액 1 (기존 결과 X열)': 2,
                    '차감금액 2 (기존 결과 AS열)': 1}
        for select in app.selectbox:
            if select.label in required:
                select.set_value(required[select.label])
        app.run()
        assert not app.exception
        run = next(button for button in app.button if '데이터 통합' in button.label)
        assert not run.disabled
        run.click().run()
        assert not app.exception
        output = app.session_state['salary_result'][1][0]
        assert read_book(output).active['V2'].value == 1000
        zipped = zipfile.ZipFile(io.BytesIO(app.session_state['ot_notice_result'][1]))
        names = zipped.namelist()
        assert len(names) == 2 and '급여DB_최종완료.xlsx' in names
        ot_output = read_book(zipped.read(next(n for n in names if n.startswith('안내문추가_'))))
        assert ot_output.active['I2'].value == '연장OT:16H'


def test_ot_missing_amount_is_not_zero():
    data, mapping, sources = fixture()
    wb = read_book(sources[0][0])
    wb.active['A2'] = None
    sources[0] = (save(wb),) + sources[0][1:]
    with pytest.raises(InputError, match='금액 또는 수식'):
        integrate(data, '급여', 2, mapping, sources)


def test_legacy_xls_embedded_theme_zip_is_not_xlsx():
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, 'w') as zipped:
        zipped.writestr('theme/theme/theme1.xml', '<theme/>')
    data = bytes.fromhex('d0cf11e0a1b11ae1') + archive.getvalue()
    assert zipfile.is_zipfile(io.BytesIO(data))
    with patch('excel_engine.shutil.which', return_value=None):
        with pytest.raises(InputError, match='변환기가 설치'):
            to_xlsx(data, 'erp.xls')


@pytest.mark.parametrize('merged', [False, True])
def test_realistic_multirow_ot_headers(merged):
    from excel_engine import detect_header_end
    wb = openpyxl.Workbook()
    ws = wb.active
    ws['A1'] = 'T/C 연장근로수당'
    for col, text in [(5, '성명'), (7, '입사일자'), (10, '연장근로수당')]:
        ws.cell(4, col, text)
    for col, text in [(10, '조출, 점심O/T'), (12, '기본연장O/T'), (14, '야간O/T'),
                      (16, '휴일근무'), (18, '휴일O/T'), (20, '계')]:
        ws.cell(5, col, text)
    for col in range(10, 20):
        ws.cell(6, col, '시간' if col % 2 == 0 else '금액')
    ws['T7'] = 'K=B+D+F+H+J'
    if merged:
        for ref in ['E4:E7', 'G4:G7', 'J4:T4', 'J5:K5', 'L5:M5', 'N5:O5', 'P5:Q5', 'R5:S5', 'T5:T6']:
            ws.merge_cells(ref)
    ws['E8'], ws['G8'], ws['T8'] = '테스트', dt.datetime(2020, 1, 2), 0
    start = detect_header(ws)
    end = detect_header_end(ws, start)
    assert (start, end) == (4, 7)
    labels = headers(ws, start, end)
    for role, col in dict(name=5, hire=7, amount=20, early=10, extension=12,
                          night=14, holiday_days=16, holiday_hours=18).items():
        assert candidates(labels, role) == [col]


def test_erp_actual_header_aliases():
    labels = {1: '직책', 2: '지급총액', 3: '식대', 4: '연차수당'}
    assert candidates(labels, 'job') == [1]
    assert candidates(labels, 'base') == [2]
    assert candidates(labels, 'subtract1') == [3]
    assert candidates(labels, 'subtract2') == [4]


def test_legacy_statement_restores_all_sheets_and_exact_text():
    from excel_engine import legacy_statement
    wb = openpyxl.Workbook()
    for ws in [wb.active, wb.create_sheet('두번째')]:
        ws['E8'] = '가상인원'
        for col, val in [(10, 0), (12, 16), (14, 8), (16, 4), (18, 9)]:
            ws.cell(8, col, val)
        ws['BA8'] = '기존 문구'
    result = read_book(legacy_statement(save(wb)))
    for ws in result:
        assert [ws.cell(8, c).value for c in range(53, 58)] == [
            '연장OT:16H', '야간OT:8H', '휴일근무:4D', '휴일OT:9H', '조출점심저녁:0H']
        assert ws.column_dimensions['J'].hidden
        assert ws.column_dimensions['AZ'].outlineLevel == 1
        assert ws.column_dimensions['BA'].width == 15
        assert ws['E8'].value == '가상인원'


def test_legacy_statement_ui_requires_no_column_selection():
    from streamlit.testing.v1 import AppTest
    from pathlib import Path
    wb = openpyxl.Workbook()
    wb.active['E8'] = '가상인원'
    wb.active['L8'] = 16
    class Upload(io.BytesIO):
        name = 'ot.xlsx'
    upload = Upload(save(wb))
    def uploads(label, **kwargs):
        if label == '문구를 추가할 OT 파일':
            return [upload]
        return [] if kwargs.get('accept_multiple_files') else None
    with patch('streamlit.file_uploader', side_effect=uploads):
        app = AppTest.from_file(str(Path('opp.py').resolve())).run(timeout=20)
        assert not app.selectbox
        button = next(b for b in app.button if b.label == '🪄 OT 급여명세서 작업실행')
        button.click().run()
        assert not app.exception
        output = app.session_state['text_result'][2]
        assert read_book(output).active['BA8'].value == '연장OT:16H'


def test_appends_only_selected_sheet_preserves_original_cells_and_parts():
    from excel_engine import append_ot_notices
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '급여'
    ws.append(['이름', '연장OT시간', '야간OT시간', '휴일근무일수', '휴일OT시간', '조출점심저녁시간'])
    ws.append(['가상인원', 16, 8, 4, 9, 0])
    ws['J2'] = '=200*2'
    ws['A2'].font = Font(bold=True)
    wb.create_sheet('집계')['A1'] = '건드리지 않을 마지막 시트'
    data = save(wb)
    original = zipfile.ZipFile(io.BytesIO(data))
    patched = io.BytesIO()
    with zipfile.ZipFile(patched, 'w') as z:
        for info in original.infolist():
            raw = original.read(info.filename)
            if info.filename == 'xl/worksheets/sheet1.xml':
                raw = raw.replace(b'<f>200*2</f><v></v>', b'<f>200*2</f><v>400</v>')
            z.writestr(info, raw)
    data = patched.getvalue()
    mapping = dict(name=1, extension=2, night=3, holiday_days=4, holiday_hours=5, early=6)
    output = append_ot_notices([(data, '급여', 1, mapping, '운영2', 'same.xlsx')])[0][2]
    before, after = zipfile.ZipFile(io.BytesIO(data)), zipfile.ZipFile(io.BytesIO(output))
    for name in before.namelist():
        if name != 'xl/worksheets/sheet1.xml':
            assert before.read(name) == after.read(name)
    result, values = read_book(output), read_book(output, True)
    assert result.sheetnames == ['급여', '집계']
    assert result['급여']['J2'].value == '=200*2'
    assert values['급여']['J2'].value == 400
    assert result['급여']['A2'].font.bold
    assert [result['급여'].cell(2,c).value for c in range(11,16)] == [
        '연장OT:16H', '야간OT:8H', '휴일근무:4D', '휴일OT:9H', '조출점심저녁:0H']
    assert result['집계'].max_column == 1
